"""Excel export for the datalog reasonableness analysis (Cpk + GU)."""

from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.engine.analysis import (
    GuProposal,
    GuVerdict,
    ItemStat,
    Verdict,
)
from src.parsers.numeric import fmt

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_VERDICT_FILL = {
    Verdict.TIGHT: PatternFill("solid", fgColor="FFC7CE"),
    Verdict.LOOSE: PatternFill("solid", fgColor="FFEB9C"),
    Verdict.OFF_CENTRE: PatternFill("solid", fgColor="FCE4D6"),
    Verdict.REASONABLE: PatternFill("solid", fgColor="C6E0B4"),
}
_GU_FILL = {
    GuVerdict.LOOSER: PatternFill("solid", fgColor="FFC7CE"),
    GuVerdict.EQUAL: PatternFill("solid", fgColor="FFEB9C"),
    GuVerdict.TIGHTER: PatternFill("solid", fgColor="C6E0B4"),
}

_CPK_HEADERS = [
    "Test item", "Unit", "样本数", "Mean", "Sigma",
    "LSL", "USL", "Cpk", "Cpk(Lo)", "Cpk(Hi)", "偏心度", "判定",
]
_GU_HEADERS = [
    "Test item", "Unit", "样本数", "Mean", "Sigma",
    "GU LSL", "GU USL", "FT LSL", "FT USL",
    "余量 Lo", "余量 Hi", "判定",
]


def _head(ws, headers, widths):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def export_analysis(
    path: str,
    stats: List[ItemStat],
    gu: List[GuProposal],
    meta: List[tuple],
    cpk_summary: dict,
    gu_summary: dict,
) -> str:
    wb = Workbook()

    ws = wb.active
    ws.title = "摘要"
    ws.append(["项", "值"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    for key, value in meta:
        ws.append([key, value])
    ws.append([])
    ws.append(["Cpk 判定分布", ""])
    for key, value in sorted(cpk_summary.items(), key=lambda kv: -kv[1]):
        ws.append([key, value])
    ws.append([])
    ws.append(["GU 判定分布", ""])
    for key, value in sorted(gu_summary.items(), key=lambda kv: -kv[1]):
        ws.append([key, value])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 46

    ws = wb.create_sheet("Cpk 分析")
    _head(ws, _CPK_HEADERS, [32, 8, 8, 14, 12, 12, 12, 10, 10, 10, 10, 10])
    for r, stat in enumerate(stats, start=2):
        values = [
            stat.name, stat.unit, stat.n,
            fmt(stat.mean, 6), fmt(stat.sigma, 6),
            fmt(stat.lo), fmt(stat.hi),
            fmt(stat.cpk, 3), fmt(stat.cpk_lo, 3), fmt(stat.cpk_hi, 3),
            fmt(stat.offset, 3), stat.verdict.value,
        ]
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
        fill = _VERDICT_FILL.get(stat.verdict)
        if fill:
            ws.cell(row=r, column=12).fill = fill
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("GU 建议")
    _head(ws, _GU_HEADERS, [32, 8, 8, 14, 12, 14, 14, 12, 12, 12, 12, 16])
    for r, item in enumerate(gu, start=2):
        values = [
            item.name, item.unit, item.n,
            fmt(item.mean, 6), fmt(item.sigma, 6),
            fmt(item.gu_lo, 6), fmt(item.gu_hi, 6),
            fmt(item.ft_lo), fmt(item.ft_hi),
            fmt(item.margin_lo, 6), fmt(item.margin_hi, 6),
            item.verdict.value,
        ]
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
        fill = _GU_FILL.get(item.verdict)
        if fill:
            ws.cell(row=r, column=12).fill = fill
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path
