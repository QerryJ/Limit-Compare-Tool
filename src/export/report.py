"""Excel export for the scenario-2 "limits to update" report.

Sheets:
  * 摘要        -- counts and the source files used
  * 待更新      -- the actionable list (what an ATE engineer actually edits)
  * 程序缺项    -- PE defines a test the program does not have
  * 基线漂移    -- PE "old program" value != program current value
  * 未匹配待确认 -- ambiguous / unmatched rows for the human to map
  * 全部明细    -- every aligned row, including unchanged ones
"""

from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.engine.align import Match, MatchStatus
from src.engine.coverage import GuVsFtReport, GuVsFtStatus
from src.engine.pe_compare import (
    BaselineRow,
    Direction,
    PeCompareResult,
    UpdateRow,
    UpdateStatus,
)
from src.model import ChangeDirection, CompareReport, CoverageReport, CoverageStatus
from src.parsers.numeric import fmt

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_FILLS = {
    Direction.LOOSEN: PatternFill("solid", fgColor="FFC7CE"),
    Direction.MIXED: PatternFill("solid", fgColor="FCE4D6"),
    Direction.TIGHTEN: PatternFill("solid", fgColor="FFEB9C"),
    Direction.FILL: PatternFill("solid", fgColor="FFF2CC"),
}
_MISSING_FILL = PatternFill("solid", fgColor="F8CBAD")
_DRIFT_FILL = PatternFill("solid", fgColor="FFC7CE")

_UPDATE_HEADERS = [
    "Function", "Test item", "Unit", "温区", "侧别",
    "当前 LSL", "当前 HSL", "PE 目标 LSL", "PE 目标 HSL",
    "Δ LSL", "Δ HSL", "方向", "状态", "PE 文件", "PE 行号", "备注",
]


def _style_header(ws, headers: List[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws, widths: List[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _update_values(row: UpdateRow) -> List:
    return [
        row.function, row.test_item, row.unit, row.zone, row.side,
        fmt(row.cur_lo), fmt(row.cur_hi), fmt(row.tgt_lo), fmt(row.tgt_hi),
        fmt(row.d_lo), fmt(row.d_hi),
        row.direction.value, row.status.value, row.pe_file, row.pe_row, row.note,
    ]


def _write_updates(ws, rows: List[UpdateRow]) -> None:
    _style_header(ws, _UPDATE_HEADERS)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(_update_values(row), start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
        if row.status is UpdateStatus.MISSING_IN_PROGRAM:
            fill = _MISSING_FILL
        else:
            fill = _FILLS.get(row.direction)
        if fill is not None:
            ws.cell(row=r_idx, column=12).fill = fill
    _autosize(ws, [18, 30, 8, 8, 7, 12, 12, 13, 13, 11, 11, 8, 12, 34, 8, 34])
    ws.auto_filter.ref = ws.dimensions


def export_pe_report(
    path: str,
    result: PeCompareResult,
    matches: List[Match],
    program_name: str = "",
    pe_files: List[str] = None,
    extra_meta: List[tuple] = None,
) -> str:
    wb = Workbook()

    ws = wb.active
    ws.title = "摘要"
    ws.append(["项", "数量"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    for key, value in (extra_meta or []):
        ws.append([key, value])
    ws.append(["对比程序", program_name])
    ws.append(["PE 文件数", len(pe_files or [])])
    for key, value in result.summary.items():
        ws.append([key, value])
    _autosize(ws, [22, 16])

    changed = [r for r in result.updates if r.status is UpdateStatus.CHANGE]
    _write_updates(wb.create_sheet("待更新"), changed)

    missing = [
        r for r in result.updates if r.status is UpdateStatus.MISSING_IN_PROGRAM
    ]
    _write_updates(wb.create_sheet("程序缺项"), missing)

    ws = wb.create_sheet("基线漂移")
    _style_header(ws, [
        "Function", "Test item", "Unit",
        "PE 记录旧值 LSL", "PE 记录旧值 HSL",
        "程序当前 LSL", "程序当前 HSL", "是否漂移", "PE 文件", "PE 行号",
    ])
    drift_rows = [b for b in result.baseline if b.drift]
    for r_idx, b in enumerate(drift_rows, start=2):
        values = [
            b.function, b.test_item, b.unit,
            fmt(b.pe_old_lo), fmt(b.pe_old_hi),
            fmt(b.prog_lo), fmt(b.prog_hi), "是", b.pe_file, b.pe_row,
        ]
        for c_idx, value in enumerate(values, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
        ws.cell(row=r_idx, column=8).fill = _DRIFT_FILL
    _autosize(ws, [18, 30, 8, 16, 16, 14, 14, 10, 34, 8])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("未匹配待确认")
    _style_header(ws, ["PE Test item", "Unit", "状态", "候选程序测试项"])
    for r_idx, m in enumerate(result.unmatched, start=2):
        status = "多候选" if m.status is MatchStatus.AMBIGUOUS else "无候选"
        values = [
            m.pe_item.name, m.pe_item.unit, status,
            " | ".join(m.candidates) if m.candidates else "",
        ]
        for c_idx, value in enumerate(values, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [34, 8, 10, 70])
    ws.auto_filter.ref = ws.dimensions

    _write_updates(wb.create_sheet("全部明细"), result.updates)

    wb.save(path)
    return path


_GUFT_HEADERS = [
    "Function", "Test item", "Unit", "温区",
    "FT LSL", "FT HSL", "GU LSL", "GU HSL", "判定", "说明",
]


def export_gu_vs_ft(
    path: str,
    report: GuVsFtReport,
    ft_name: str = "",
    gu_name: str = "",
) -> str:
    """Excel export for scenario 4: GU limit must sit strictly inside FT."""
    wb = Workbook()

    ws = wb.active
    ws.title = "摘要"
    ws.append(["项", "数量"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    ws.append(["FT 程序", ft_name])
    ws.append(["GU 程序", gu_name])
    for key, value in report.summary.items():
        ws.append([key, value])
    _autosize(ws, [26, 16])

    failed = [r for r in report.items if r.status is GuVsFtStatus.OUTSIDE]
    ws = wb.create_sheet("GU 比 FT 宽(失效)")
    _style_header(ws, _GUFT_HEADERS)
    for r_idx, r in enumerate(failed, start=2):
        for z in r.zones:
            if z.status is not GuVsFtStatus.OUTSIDE:
                continue
            values = [
                r.item.function, r.item.test_item, r.item.unit, z.zone.value,
                fmt(z.ft.lo) if z.ft else "", fmt(z.ft.hi) if z.ft else "",
                fmt(z.gu.lo) if z.gu else "", fmt(z.gu.hi) if z.gu else "",
                z.status.value, z.detail,
            ]
            for c_idx, value in enumerate(values, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
            ws.cell(row=r_idx, column=9).fill = _MISSING_FILL
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 16, 60])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("全部明细")
    _style_header(ws, _GUFT_HEADERS)
    for r_idx, r in enumerate(report.items, start=2):
        for z in r.zones:
            values = [
                r.item.function, r.item.test_item, r.item.unit, z.zone.value,
                fmt(z.ft.lo) if z.ft else "", fmt(z.ft.hi) if z.ft else "",
                fmt(z.gu.lo) if z.gu else "", fmt(z.gu.hi) if z.gu else "",
                z.status.value, z.detail,
            ]
            for c_idx, value in enumerate(values, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 16, 60])
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# FV-1: FT vs QA coverage report
# ---------------------------------------------------------------------------
_COVERAGE_HEADERS = [
    "Function", "Test item", "Unit", "温区",
    "FT LSL", "FT HSL", "QA LSL", "QA HSL", "低边裕度", "高边裕度", "判定", "说明",
]


def _coverage_row(entry, z) -> List:
    return [
        entry.item.function, entry.item.test_item, entry.item.unit, z.zone.value,
        fmt(z.ft.lo) if z.ft else "", fmt(z.ft.hi) if z.ft else "",
        fmt(z.qa.lo) if z.qa else "", fmt(z.qa.hi) if z.qa else "",
        fmt(z.gb_lo) if z.gb_lo is not None else "",
        fmt(z.gb_hi) if z.gb_hi is not None else "",
        z.status.value, z.detail,
    ]


def export_coverage_report(
    path: str,
    report: CoverageReport,
    ft_name: str = "",
    qa_name: str = "",
) -> str:
    """Excel export for FV-1: FT-vs-QA coverage inside one program.

    Sheets: 摘要 / 未覆盖(需处理) / 局部缺失 / 整体缺失 / 全部明细.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "摘要"
    ws.append(["项", "数量"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    ws.append(["FT 程序", ft_name])
    ws.append(["QA 程序", qa_name])
    for key, value in report.summary.items():
        ws.append([key, value])
    _autosize(ws, [22, 16])

    failed = [e for e in report.items
              if e.status in (CoverageStatus.UNCOVERED, CoverageStatus.ERROR)]
    ws = wb.create_sheet("未覆盖(需处理)")
    _style_header(ws, _COVERAGE_HEADERS)
    for r_idx, e in enumerate(failed, start=2):
        for z in e.zones:
            if z.status is CoverageStatus.COVERED:
                continue
            for c_idx, value in enumerate(_coverage_row(e, z), start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
            ws.cell(row=r_idx, column=11).fill = _MISSING_FILL
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 10, 10, 12, 60])
    ws.auto_filter.ref = ws.dimensions

    partial = [e for e in report.items if e.status is CoverageStatus.PARTIAL]
    ws = wb.create_sheet("局部缺失")
    _style_header(ws, _COVERAGE_HEADERS)
    for r_idx, e in enumerate(partial, start=2):
        for z in e.zones:
            for c_idx, value in enumerate(_coverage_row(e, z), start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 10, 10, 12, 60])
    ws.auto_filter.ref = ws.dimensions

    missing = [e for e in report.items if e.status is CoverageStatus.MISSING]
    ws = wb.create_sheet("整体缺失")
    _style_header(ws, _COVERAGE_HEADERS)
    for r_idx, e in enumerate(missing, start=2):
        for z in e.zones:
            for c_idx, value in enumerate(_coverage_row(e, z), start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 10, 10, 12, 60])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("全部明细")
    _style_header(ws, _COVERAGE_HEADERS)
    for r_idx, e in enumerate(report.items, start=2):
        for z in e.zones:
            for c_idx, value in enumerate(_coverage_row(e, z), start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 8, 12, 12, 12, 12, 10, 10, 12, 60])
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# FV-4: before-vs-after change review
# ---------------------------------------------------------------------------
_CHANGE_DETAIL_HEADERS = [
    "Function", "Test item", "Unit", "温区", "侧别",
    "旧 LSL", "旧 HSL", "新 LSL", "新 HSL", "Δ LSL", "Δ HSL", "方向",
]
_CHANGE_SUMMARY_HEADERS = ["Function", "Test item", "Unit", "方向", "变更说明"]


def _change_cell_row(d, c) -> List:
    return [
        d.function, d.test_item, d.unit, c.zone.value, c.side.value,
        fmt(c.a_lo) if c.a_lo is not None else "",
        fmt(c.a_hi) if c.a_hi is not None else "",
        fmt(c.b_lo) if c.b_lo is not None else "",
        fmt(c.b_hi) if c.b_hi is not None else "",
        fmt(c.delta_lo) if c.delta_lo is not None else "",
        fmt(c.delta_hi) if c.delta_hi is not None else "",
        d.direction.value,
    ]


def _change_summary_row(d) -> List:
    return [d.function, d.test_item, d.unit, d.direction.value, d.change_detail]


def export_compare_report(
    path: str,
    report: CompareReport,
    base_name: str = "",
    new_name: str = "",
) -> str:
    """Excel export for FV-4: before-vs-after program change review.

    Sheets: 摘要 / 放宽风险 / 新增与移除 / 变更明细(逐格) / 全部明细.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "摘要"
    ws.append(["项", "数量"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    ws.append(["基准程序", base_name])
    ws.append(["新程序", new_name])
    for key, value in report.summary.items():
        ws.append([key, value])
    _autosize(ws, [22, 16])

    risky = [d for d in report.items
             if d.direction in (ChangeDirection.LOOSER, ChangeDirection.MIXED)]
    ws = wb.create_sheet("放宽风险")
    _style_header(ws, _CHANGE_SUMMARY_HEADERS)
    for r_idx, d in enumerate(risky, start=2):
        for c_idx, value in enumerate(_change_summary_row(d), start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
        ws.cell(row=r_idx, column=4).fill = _MISSING_FILL
    _autosize(ws, [18, 32, 8, 10, 80])
    ws.auto_filter.ref = ws.dimensions

    added_removed = [d for d in report.items
                     if d.direction in (ChangeDirection.ONLY_IN_B, ChangeDirection.ONLY_IN_A)]
    ws = wb.create_sheet("新增与移除")
    _style_header(ws, _CHANGE_SUMMARY_HEADERS)
    for r_idx, d in enumerate(added_removed, start=2):
        for c_idx, value in enumerate(_change_summary_row(d), start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 10, 80])
    ws.auto_filter.ref = ws.dimensions

    changed = [
        d for d in report.items
        if d.direction not in (
            ChangeDirection.SAME, ChangeDirection.ONLY_IN_A, ChangeDirection.ONLY_IN_B
        ) and d.zones
    ]
    ws = wb.create_sheet("变更明细(逐格)")
    _style_header(ws, _CHANGE_DETAIL_HEADERS)
    for r_idx, d in enumerate(changed, start=2):
        for c in d.zones:
            for c_idx, value in enumerate(_change_cell_row(d, c), start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
            if d.direction in (ChangeDirection.LOOSER, ChangeDirection.MIXED):
                ws.cell(row=r_idx, column=12).fill = _MISSING_FILL
    _autosize(ws, [18, 30, 8, 8, 7, 12, 12, 12, 12, 11, 11, 10])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("全部明细")
    _style_header(ws, _CHANGE_SUMMARY_HEADERS)
    for r_idx, d in enumerate(report.items, start=2):
        for c_idx, value in enumerate(_change_summary_row(d), start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    _autosize(ws, [18, 32, 8, 10, 80])
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path
