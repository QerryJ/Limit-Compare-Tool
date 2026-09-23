"""Audit metadata for exported reports.

Every report must answer "when, with which tool version, on which inputs, with
which parameters" -- otherwise it cannot be used as audit evidence. The
dedicated exporters own their own sheets, so this module *post-processes* the
finished workbook and prepends a ``稽核信息`` sheet; that keeps one code path
for all six modes instead of touching every exporter.
"""

from typing import List, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.ui.theme import hex_color

SHEET_NAME = "稽核信息"


def _argb(value: str) -> str:
    """``#rrggbb`` -> openpyxl 需要的 aRGB ``FFrrggbb``。"""
    raw = value.lstrip("#")
    return raw if len(raw) == 8 else "FF" + raw


def write_audit_sheet(
    path: str,
    info: List[Tuple[str, str]],
    summary: str = "",
) -> str:
    """Prepend an audit sheet to an already-written XLSX report."""
    wb = load_workbook(path)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME, 0)

    title = ws.cell(row=1, column=1, value=SHEET_NAME)
    title.font = Font(bold=True, size=13, color=_argb(hex_color("--lc-text")))

    row = 3
    if summary:
        head = ws.cell(row=row, column=1, value="对比总结")
        head.font = Font(bold=True)
        head.fill = PatternFill("solid", fgColor=_argb(hex_color("--lc-surface")))
        cell = ws.cell(row=row, column=2, value=summary)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 2

    key_font = Font(bold=True)
    key_fill = PatternFill("solid", fgColor=_argb(hex_color("--lc-surface")))
    for key, value in info:
        key_cell = ws.cell(row=row, column=1, value=str(key))
        key_cell.font = key_font
        key_cell.fill = key_fill
        key_cell.alignment = Alignment(vertical="top")
        val_cell = ws.cell(row=row, column=2, value=str(value))
        val_cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 96
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def enable_header_filters(path: str, skip=(SHEET_NAME,)) -> str:
    """给导出报告的**数据表**开启「表头筛选 + 冻结首行」。

    工程师拿到 XLSX 的第一件事就是点判定列、只看未通过的项；
    没有 AutoFilter 就得手工套筛选，容易漏。稽核/摘要这类键值表跳过。
    """
    wb = load_workbook(path)
    changed = False
    for ws in wb.worksheets:
        # 键值型表（稽核信息 / 摘要）非二维明细，加筛选反而干扰阅读
        if ws.title in skip or "摘要" in ws.title:
            continue
        if ws.max_row < 2 or ws.max_column < 3:
            continue
        if not ws.freeze_panes:
            ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        changed = True
    if changed:
        wb.save(path)
    return path
