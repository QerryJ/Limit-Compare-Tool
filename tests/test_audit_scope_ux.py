"""v0.4.0: audit metadata, export scope, search & sort, stylesheet from tokens.

Headless (``QT_QPA_PLATFORM=offscreen``). Skipped where PySide6 is absent.
"""

import os
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402

pytest.importorskip("PySide6", reason="PySide6 not installed in this environment")

from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

from src.export.audit import SHEET_NAME, enable_header_filters, write_audit_sheet  # noqa: E402
from src.export.html_report import export_html_report  # noqa: E402
from src.ui.main_window import APP_AUTHOR, APP_VERSION, MainWindow, _SortItem  # noqa: E402
from src.ui.style import build_stylesheet  # noqa: E402
from src.ui.theme import hex_color  # noqa: E402


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def _seed_rows(window, n=3):
    """给窗口塞入可控的结果行，避免依赖真实素材。"""
    window._rows = [[f"FUNC_{i}", f"ITEM_{i}", "过紧" if i == 1 else "合理"]
                    for i in range(n)]
    window._headers = ["Function", "Test item", "判定"]
    window._verdict_col = 2
    window._anomaly_set = {"过紧"}
    return window._rows


# --------------------------------------------------------------- GR-3 search
def test_search_narrows_to_matching_rows(app):
    window = MainWindow()
    _seed_rows(window)
    window.search_edit.setText("func_1")
    window._apply_filter()
    assert window.table.rowCount() == 1
    assert window.table.item(0, 0).text() == "FUNC_1"


def test_search_combines_with_verdict_filter(app):
    window = MainWindow()
    _seed_rows(window)
    window.search_edit.setText("func_1")     # 命中 FUNC_1（判定=过紧）
    window.filter_combo.setCurrentText("仅异常")
    window._apply_filter()
    assert window.table.rowCount() == 1      # 搜索与筛选叠加生效

    window.search_edit.setText("func_0")     # FUNC_0 判定=合理，不在异常集
    window._apply_filter()
    assert window.table.rowCount() == 0


def test_clear_files_resets_search(app):
    window = MainWindow()
    window.search_edit.setText("abc")
    window.clear_btn.click()
    assert window.search_edit.text() == ""


# ----------------------------------------------------------------- GR-3 sort
def test_sort_item_orders_numerically_not_lexically():
    assert _SortItem("9") < _SortItem("10")      # 字符串排序会错，数值排序正确
    assert not (_SortItem("10") < _SortItem("9"))


def test_sort_item_falls_back_to_text():
    assert _SortItem("N/A") < _SortItem("mV")


# --------------------------------------------------------------- GR-1 audit
def test_audit_info_keeps_time_filter_and_sources(app):
    """按用户要求：稽核信息只保留 导出时间 / 筛选信息 / 源文件全路径。"""
    window = MainWindow()
    window.program_edit.setText("C:/tmp/FT.pds")
    window.datalog_edit.setText("C:/tmp/log.csv")
    info = dict(window._audit_info(10, "全量"))
    assert info["导出时间"]
    assert info["结果筛选"] == "全部"
    assert info["导出口径"] == "全量"
    assert info["导出行数"] == "10"
    assert info["程序 / 基准"] == "C:/tmp/FT.pds"
    assert info["datalog"] == "C:/tmp/log.csv"


def test_audit_info_excludes_tool_version_and_author(app):
    window = MainWindow()
    info = dict(window._audit_info(1, "全量"))
    blob = " ".join(info.keys()) + " " + " ".join(map(str, info.values()))
    assert "工具" not in info
    assert "作者" not in info
    assert APP_AUTHOR not in blob
    assert APP_VERSION not in blob


def test_audit_info_excludes_mode_and_sigma(app):
    from src.ui.main_window import MODE_ANALYSIS

    window = MainWindow()
    window.mode_combo.setCurrentText(MODE_ANALYSIS)
    info = dict(window._audit_info(1, "全量"))
    assert "GU σ 倍数" not in info
    assert "对比模式" not in info
    assert "温区" not in info


def test_write_audit_sheet_prepends_sheet(tmp_path):
    from openpyxl import Workbook, load_workbook

    path = str(tmp_path / "report.xlsx")
    wb = Workbook()
    wb.active.title = "摘要"
    wb.active["A1"] = "x"
    wb.save(path)

    write_audit_sheet(path, [("工具", "Limit Compare v9.9"), ("温区", "Room")])

    wb2 = load_workbook(path)
    assert wb2.sheetnames[0] == SHEET_NAME
    sheet = wb2[SHEET_NAME]
    assert sheet["A3"].value == "工具"
    assert sheet["B3"].value == "Limit Compare v9.9"
    assert "摘要" in wb2.sheetnames


def test_html_report_embeds_audit_block(tmp_path):
    path = str(tmp_path / "r.html")
    export_html_report(
        path, "T", ["项", "判定"], [["A", "过紧"]],
        audit=[("工具", "Limit Compare v0.4.0"), ("导出口径", "全量")],
    )
    text = open(path, encoding="utf-8").read()
    assert "稽核信息" in text
    assert "Limit Compare v0.4.0" in text
    assert "导出口径" in text
    assert "<link" not in text and "cdn" not in text.lower()   # 仍是离线单文件


def test_html_report_has_header_filter_and_summary(tmp_path):
    """导出的报告也要能直接筛出状态列未通过的项。"""
    path = str(tmp_path / "r.html")
    export_html_report(
        path, "T", ["项", "判定"], [["A", "未覆盖"], ["B", "覆盖"]],
        summary="未覆盖 1 项", audit=[("导出时间", "2026-09-24 00:00:00")],
    )
    text = open(path, encoding="utf-8").read()
    assert 'id="lc-vfilter"' in text          # 判定列下拉
    assert ">未覆盖<" in text and ">覆盖<" in text
    assert 'class="data"' in text             # 只筛数据表，不筛稽核表
    assert "对比总结" in text and "未覆盖 1 项" in text
    assert "<link" not in text and "cdn" not in text.lower()


def test_enable_header_filters_only_on_data_sheets(tmp_path):
    from openpyxl import Workbook, load_workbook

    path = str(tmp_path / "r.xlsx")
    wb = Workbook()
    summary = wb.active
    summary.title = "摘要"                     # 键值表，不该加筛选
    summary.append(["项", "值"])
    detail = wb.create_sheet("全部明细")
    detail.append(["Function", "Test item", "判定"])
    detail.append(["F1", "T1", "未覆盖"])
    detail.append(["F2", "T2", "覆盖"])
    wb.save(path)

    enable_header_filters(path)

    wb2 = load_workbook(path)
    assert wb2["摘要"].auto_filter.ref is None
    assert wb2["全部明细"].auto_filter.ref == "A1:C3"
    assert wb2["全部明细"].freeze_panes == "A2"


# --------------------------------------------------------------- GR-2 scope
_CAPTURED = []


def _fake_add_button(self, *args, **kwargs):
    btn = object()
    _CAPTURED.append(btn)
    return btn


def _call_scope(window, pick):
    """执行 _resolve_export_scope，pick=0 全量 / 1 当前视图 / 2 取消。"""
    _CAPTURED.clear()
    with mock.patch.object(QMessageBox, "addButton", _fake_add_button), \
         mock.patch.object(QMessageBox, "exec", lambda self: 0), \
         mock.patch.object(QMessageBox, "clickedButton",
                           lambda self: _CAPTURED[pick] if pick < len(_CAPTURED) else None):
        return window._resolve_export_scope()


def test_export_scope_no_dialog_when_unfiltered(app):
    window = MainWindow()
    _seed_rows(window)
    window._apply_filter()                      # 筛选=全部
    rows, scope = window._resolve_export_scope()
    assert len(rows) == 3
    assert "全部" in scope


def test_export_scope_can_export_full(app):
    window = MainWindow()
    _seed_rows(window)
    window.filter_combo.setCurrentText("仅异常")
    window._apply_filter()
    assert window.table.rowCount() == 1         # 视图只剩 1 行，全量 3 行
    rows, scope = _call_scope(window, 0)
    assert len(rows) == 3
    assert scope == "全量"


def test_export_scope_can_export_current_view(app):
    window = MainWindow()
    _seed_rows(window)
    window.filter_combo.setCurrentText("仅异常")
    window._apply_filter()
    rows, scope = _call_scope(window, 1)
    assert len(rows) == 1
    assert "当前视图" in scope


def test_export_scope_cancel_returns_none(app):
    window = MainWindow()
    _seed_rows(window)
    window.filter_combo.setCurrentText("仅异常")
    window._apply_filter()
    rows, scope = _call_scope(window, 2)
    assert rows is None and scope == ""


# ------------------------------------------------------------ GR-4 messaging
def test_missing_program_hint_names_the_missing_side(app):
    window = MainWindow()
    hint = window._missing_files_hint("FT 程序", "GU 程序")
    assert "FT 程序" in hint and "GU 程序" in hint
    assert "拖" in hint                          # 给出可操作的补救方式


# ----------------------------------------------------------------- UI style
def test_stylesheet_is_generated_from_tokens():
    css = build_stylesheet("comfortable")
    assert "QPushButton#primary" in css
    assert hex_color("--lc-accent").lstrip("#") in css
    assert build_stylesheet("compact") != build_stylesheet("comfortable")


def test_window_applies_stylesheet(app):
    window = MainWindow()
    assert window.styleSheet()
    assert window.run_btn.objectName() == "primary"
    assert window.author_label.objectName() == "author"
