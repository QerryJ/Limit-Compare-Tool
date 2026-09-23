"""FV-1 coverage and FV-4 change-review engines."""

import pytest

from src.engine.compare import compare_programs, risky_changes
from src.engine.coverage import evaluate_coverage, zone_coverage
from src.export.report import export_compare_report, export_coverage_report
from src.model import (
    Bound,
    ChangeDirection,
    CoverageStatus,
    LimitItem,
    MachineType,
    ProgramModel,
    SourceType,
    TempZone,
)


def _item(name, ft, qa, function="IQ") -> LimitItem:
    item = LimitItem(function=function, test_item=name, unit="V",
                     machine_type=MachineType.ETS88)
    if ft is not None:
        item.set_bound(TempZone.ROOM, SourceType.FT, Bound(lo=ft[0], hi=ft[1]))
    if qa is not None:
        item.set_bound(TempZone.ROOM, SourceType.QA, Bound(lo=qa[0], hi=qa[1]))
    return item


def _program(*items) -> ProgramModel:
    return ProgramModel(id="p", name="p", machine_type=MachineType.ETS88,
                        functions=["IQ"], items=list(items))


# ------------------------------------------------------------------ coverage
def test_qa_containing_ft_is_covered():
    result = zone_coverage(
        TempZone.ROOM, Bound(lo=2.0, hi=3.0), Bound(lo=1.5, hi=3.5)
    )
    assert result.status is CoverageStatus.COVERED
    assert result.gb_lo == 0.5 and result.gb_hi == 0.5


def test_qa_narrower_than_ft_is_uncovered_with_actionable_detail():
    result = zone_coverage(
        TempZone.ROOM, Bound(lo=2.0, hi=3.0), Bound(lo=2.2, hi=2.8)
    )
    assert result.status is CoverageStatus.UNCOVERED
    assert result.gb_lo is not None and result.gb_lo < 0
    assert "低边未覆盖" in result.detail and "高边未覆盖" in result.detail


def test_reversed_ft_interval_is_error_not_uncovered():
    result = zone_coverage(
        TempZone.ROOM, Bound(lo=3.0, hi=2.0), Bound(lo=1.0, hi=4.0)
    )
    assert result.status is CoverageStatus.ERROR
    assert "反向" in result.detail


def test_missing_side_is_missing():
    assert zone_coverage(TempZone.ROOM, Bound(lo=1.0, hi=2.0), None).status is CoverageStatus.MISSING


def test_coverage_report_summary():
    program = _program(
        _item("A", (2.0, 3.0), (1.5, 3.5)),
        _item("B", (2.0, 3.0), (2.2, 2.8)),
    )
    report = evaluate_coverage(program, [TempZone.ROOM])
    assert report.summary[CoverageStatus.COVERED.value] == 1
    assert report.summary[CoverageStatus.UNCOVERED.value] == 1
    assert report.summary["覆盖率"] == 50.0


# ------------------------------------------------------------------- compare
def test_tighten_and_loosen_direction():
    base = _program(_item("A", (2.0, 3.0), (1.5, 3.5)))
    new = _program(_item("A", (2.2, 2.8), (1.8, 3.2)))
    report = compare_programs(base, new, [TempZone.ROOM])
    assert report.items[0].direction is ChangeDirection.TIGHTER
    assert risky_changes(report) == []

    loosened = compare_programs(new, base, [TempZone.ROOM])
    assert loosened.items[0].direction is ChangeDirection.LOOSER
    assert len(risky_changes(loosened)) == 1


def test_mixed_change_counts_as_risk():
    base = _program(_item("A", (2.0, 3.0), None))
    new = _program(_item("A", (2.2, 3.5), None))
    report = compare_programs(base, new, [TempZone.ROOM])
    assert report.items[0].direction is ChangeDirection.MIXED
    assert report.summary["放宽风险"] >= 1
    assert len(risky_changes(report)) == 1


def test_added_and_removed_items():
    base = _program(_item("OLD", (1.0, 2.0), None))
    new = _program(_item("NEW", (1.0, 2.0), None))
    report = compare_programs(base, new, [TempZone.ROOM])
    directions = {d.test_item: d.direction for d in report.items}
    assert directions["NEW"] is ChangeDirection.ONLY_IN_B
    assert directions["OLD"] is ChangeDirection.ONLY_IN_A


def test_identical_programs_produce_no_changes():
    base = _program(_item("A", (2.0, 3.0), (1.5, 3.5)))
    new = _program(_item("A", (2.0, 3.0), (1.5, 3.5)))
    report = compare_programs(base, new, [TempZone.ROOM])
    assert report.items[0].direction is ChangeDirection.SAME
    assert report.summary.get("放宽风险", 0) == 0


def test_floating_point_noise_is_not_a_change():
    base = _program(_item("A", (2.0, 3.0), None))
    new = _program(_item("A", (2.0 + 1e-13, 3.0), None))
    report = compare_programs(base, new, [TempZone.ROOM])
    assert report.items[0].direction is ChangeDirection.SAME


# ------------------------------------------------------------------- exporters
def test_export_coverage_report_sheets(tmp_path):
    program = _program(
        _item("A", (2.0, 3.0), (1.5, 3.5)),
        _item("B", (2.0, 3.0), (2.2, 2.8)),
        _item("C", (2.0, 3.0), None),  # 整体缺失（无 QA 侧）
    )
    report = evaluate_coverage(program, [TempZone.ROOM])
    out = tmp_path / "cov.xlsx"
    export_coverage_report(str(out), report, ft_name="ft", qa_name="qa")

    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    assert wb.sheetnames == [
        "摘要", "未覆盖(需处理)", "局部缺失", "整体缺失", "全部明细",
    ]
    summary = {r[0].value: r[1].value for r in wb["摘要"].iter_rows(min_row=2)}
    # 可比温区 = A + B = 2；覆盖 = A = 1 → 50%
    assert summary["覆盖率"] == 50.0
    assert summary["整体缺失"] == 1


def test_export_compare_report_sheets(tmp_path):
    base = _program(_item("A", (2.0, 3.0), None))
    new = _program(_item("A", (1.8, 3.2), None))  # Lo/Hi 均放宽 -> 漏放风险
    report = compare_programs(base, new, [TempZone.ROOM])
    out = tmp_path / "cmp.xlsx"
    export_compare_report(str(out), report, base_name="b", new_name="n")

    from openpyxl import load_workbook

    wb = load_workbook(str(out))
    assert wb.sheetnames == [
        "摘要", "放宽风险", "新增与移除", "变更明细(逐格)", "全部明细",
    ]
    risky = list(wb["放宽风险"].iter_rows(min_row=2))
    assert len(risky) == 1  # A 放宽（Lo 2.0→1.8, Hi 3.0→3.2）
